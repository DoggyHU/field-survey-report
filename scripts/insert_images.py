#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 DOCX 里的 @@IMG:文件名|宽cm@@ 占位替换成实际图片（field-survey-report skill）

背景：gongwen-docx 只做排版、不内嵌图片，所以正文先写占位标记，转完 DOCX 再用本脚本插图。

用法:
    python insert_images.py --docx 报告.docx --base-dir .

约定:
- 占位标记单独成段：@@IMG:路线图_4566590.png|15.6@@
- 相对路径按 --base-dir 解析
- 图片段落居中、取消首行缩进、行距改单倍；并设 keep_with_next（避免图题跑到下一页）
- 以「图N …」开头的图题会清掉 keep_with_next，保证题注紧跟其图
"""
import argparse, os, re

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

MARK = re.compile(r"@@IMG:(.+?)\|([0-9.]+)@@")
FIGCAP = re.compile(r"^图\s*[0-9A-Z一二三四五六七八九十]+")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", required=True)
    ap.add_argument("--base-dir", default=".")
    a = ap.parse_args()

    doc = Document(a.docx)
    n = 0
    for p in doc.paragraphs:
        t = p.text.strip()
        m = MARK.fullmatch(t)
        if m:
            fn, w = m.group(1).strip(), float(m.group(2))
            path = fn if os.path.isabs(fn) else os.path.join(a.base_dir, fn)
            if not os.path.exists(path):
                print("  [warn] missing image:", path); continue
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf = p.paragraph_format
            pf.first_line_indent = 0; pf.left_indent = 0; pf.right_indent = 0
            pf.space_before = Pt(8); pf.space_after = Pt(2)
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
            pf.keep_with_next = True; pf.keep_together = True
            p.add_run().add_picture(path, width=Cm(w))
            n += 1
            continue
        if FIGCAP.match(t):
            p.paragraph_format.keep_with_next = False
            p.paragraph_format.keep_together = True
    doc.save(a.docx)
    print("inserted %d image(s) into %s" % (n, a.docx))

if __name__ == "__main__":
    main()
